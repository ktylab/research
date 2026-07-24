#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
/******************************************************************************
 * File Name : generate_layer34_benchmark_ppt.py
 * Project : infinitPINN / SorTech UI Research
 * Module : research/deliverables
 *
 * Purpose
 * ----------------------------------------------------------------------------
 * ISA-95 Layer 3~4(MES/ERP/APS) 웹 업무시스템 UI 벤치마크 PPT를 생성한다.
 *
 * Responsibilities
 * ----------------------------------------------------------------------------
 * - 한국/글로벌 제품 맵 슬라이드 구성
 * - 화면 유형 12종 상세 슬라이드 생성
 * - deliverables 및 artifacts 경로에 PPTX 저장
 *
 * Dependencies
 * ----------------------------------------------------------------------------
 * - python-pptx
 * - lxml
 *
 * Author : Cursor Agent
 * Created : 2026-07-24
 * Last Update : 2026-07-24
 *
 * ----------------------------------------------------------------------------
 * Revision History
 * ----------------------------------------------------------------------------
 * Ver   Date       Author         Description
 * ----- ---------- -------------- -------------------------------
 * 1.0   2026-07-24 Cursor Agent   Initial creation
 ******************************************************************************/
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor as RgbColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from lxml import etree

# Visual system: industrial slate + teal (avoid purple/cream AI clichés)
C_BG = RgbColor(0x0F, 0x17, 0x1F)
C_SURFACE = RgbColor(0x17, 0x24, 0x32)
C_CARD = RgbColor(0x1E, 0x2F, 0x40)
C_ACCENT = RgbColor(0x1F, 0xB5, 0xA0)  # teal
C_ACCENT2 = RgbColor(0xF0, 0xA5, 0x3A)  # amber signal
C_TEXT = RgbColor(0xF2, 0xF5, 0xF8)
C_MUTED = RgbColor(0x9A, 0xAA, 0xBA)
C_LINE = RgbColor(0x2E, 0x42, 0x56)
C_WHITE = RgbColor(0xFF, 0xFF, 0xFF)
C_LIGHT_BG = RgbColor(0xF5, 0xF7, 0xF9)
C_LIGHT_TEXT = RgbColor(0x1A, 0x2A, 0x3A)
C_LIGHT_MUTED = RgbColor(0x5A, 0x6B, 0x7C)
C_LIGHT_CARD = RgbColor(0xFF, 0xFF, 0xFF)
C_LIGHT_ACCENT = RgbColor(0x0D, 0x8F, 0x7C)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_run(run, size=14, bold=False, color=C_TEXT, font="Noto Sans KR"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    # East Asian font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if ea is None:
        ea = etree.SubElement(rPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    ea.set("typeface", font)


def add_textbox(slide, left, top, width, height, text, size=14, bold=False,
                color=C_TEXT, align=PP_ALIGN.LEFT, font="Noto Sans KR"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def add_para(tf, text, size=13, bold=False, color=C_TEXT, space_before=6, space_after=2):
    p = tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return p


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    fill_shape(shape, color)
    return shape


def round_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    fill_shape(shape, color)
    # reduce rounding
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def accent_bar(slide, left, top, width=Inches(0.08), height=Inches(0.45), color=C_ACCENT):
    return rect(slide, left, top, width, height, color)


def dark_bg(slide):
    rect(slide, 0, 0, SLIDE_W, SLIDE_H, C_BG)


def light_bg(slide):
    rect(slide, 0, 0, SLIDE_W, SLIDE_H, C_LIGHT_BG)


def footer(slide, page, total, dark=True):
    color = C_MUTED if dark else C_LIGHT_MUTED
    add_textbox(slide, Inches(0.5), Inches(7.1), Inches(8), Inches(0.3),
                "SorTech · infinitPINN UI/UX Benchmark · Confidential",
                size=10, color=color)
    add_textbox(slide, Inches(11.5), Inches(7.1), Inches(1.5), Inches(0.3),
                f"{page} / {total}", size=10, color=color, align=PP_ALIGN.RIGHT)


def section_title(slide, title, subtitle=None, dark=True):
    accent_bar(slide, Inches(0.5), Inches(0.35), height=Inches(0.5))
    tc = C_TEXT if dark else C_LIGHT_TEXT
    mc = C_MUTED if dark else C_LIGHT_MUTED
    add_textbox(slide, Inches(0.7), Inches(0.3), Inches(12), Inches(0.5),
                title, size=26, bold=True, color=tc)
    if subtitle:
        add_textbox(slide, Inches(0.7), Inches(0.85), Inches(12), Inches(0.35),
                    subtitle, size=13, color=mc)


def card_with_text(slide, left, top, width, height, title, lines, dark=True,
                   title_color=None):
    bg = C_CARD if dark else C_LIGHT_CARD
    tc = C_TEXT if dark else C_LIGHT_TEXT
    mc = C_MUTED if dark else C_LIGHT_MUTED
    if title_color is None:
        title_color = C_ACCENT if dark else C_LIGHT_ACCENT
    # light cards need a subtle border via line
    shape = round_rect(slide, left, top, width, height, bg)
    if not dark:
        shape.line.color.rgb = RgbColor(0xD5, 0xDE, 0xE7)
        shape.line.width = Pt(1)
    add_textbox(slide, left + Inches(0.18), top + Inches(0.12),
                width - Inches(0.3), Inches(0.35),
                title, size=14, bold=True, color=title_color)
    box = slide.shapes.add_textbox(
        left + Inches(0.18), top + Inches(0.45),
        width - Inches(0.3), height - Inches(0.55))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for line in lines:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = "• " + line
        set_run(run, size=11, color=mc)
    return shape


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]
    slides_meta = []

    # ---------- 1 Title ----------
    s = prs.slides.add_slide(blank)
    dark_bg(s)
    rect(s, 0, 0, Inches(0.18), SLIDE_H, C_ACCENT)
    add_textbox(s, Inches(0.7), Inches(1.8), Inches(11), Inches(0.4),
                "UI/UX BENCHMARK RESEARCH", size=14, bold=True, color=C_ACCENT)
    add_textbox(s, Inches(0.7), Inches(2.3), Inches(12), Inches(1.2),
                "ISA-95 Layer 3~4\n웹 업무시스템 화면 유형 벤치마크",
                size=36, bold=True, color=C_TEXT)
    add_textbox(s, Inches(0.7), Inches(4.0), Inches(11), Inches(0.8),
                "MES · ERP · APS 한국/글로벌 시장 조사\n소르테크 infinitPINN 플랫폼 시안 설계를 위한 자료",
                size=16, color=C_MUTED)
    add_textbox(s, Inches(0.7), Inches(5.5), Inches(11), Inches(0.4),
                "2026.07.24  |  SorTech Research", size=13, color=C_MUTED)
    slides_meta.append(("title", True))

    # ---------- 2 Agenda ----------
    s = prs.slides.add_slide(blank)
    dark_bg(s)
    section_title(s, "목차", "조사 범위와 산출물 구조")
    items = [
        ("01", "조사 범위", "ISA-95 Layer 3~4, Web UI 중심"),
        ("02", "한국 시장 맵", "MES / ERP / APS 주요 벤더·제품"),
        ("03", "글로벌 시장 맵", "Siemens, SAP, Rockwell, AVEVA 등"),
        ("04", "화면 유형 12종", "유형별 레이아웃·인터랙션·벤치마크"),
        ("05", "소르테크 시사점", "infinitPINN 차별화 방향"),
    ]
    for i, (num, t, d) in enumerate(items):
        y = Inches(1.5) + Inches(i * 0.95)
        round_rect(s, Inches(0.7), y, Inches(11.9), Inches(0.8), C_CARD)
        add_textbox(s, Inches(1.0), y + Inches(0.18), Inches(1), Inches(0.45),
                    num, size=22, bold=True, color=C_ACCENT)
        add_textbox(s, Inches(2.2), y + Inches(0.1), Inches(6), Inches(0.35),
                    t, size=18, bold=True, color=C_TEXT)
        add_textbox(s, Inches(2.2), y + Inches(0.42), Inches(9), Inches(0.3),
                    d, size=12, color=C_MUTED)
    slides_meta.append(("agenda", True))

    # ---------- 3 Scope ----------
    s = prs.slides.add_slide(blank)
    light_bg(s)
    section_title(s, "조사 범위 — ISA-95 Layer 3~4", "웹 기반 업무시스템 UI에 집중", dark=False)
    layers = [
        ("Layer 4", "경영·자원계획", "ERP / SCM / 구매·영업·재무",
         "경영진, 구매, 영업, 생산계획", C_LIGHT_ACCENT),
        ("Layer 3", "제조운영", "MES·MOM / APS / WMS / QMS",
         "감독, 오퍼레이터, 품질, 스케줄러", RgbColor(0x2B, 0x6C, 0xB0)),
        ("Layer 2↓", "제어 (참고)", "SCADA / HMI / PLC",
         "본 조사는 패턴만 참고", C_LIGHT_MUTED),
    ]
    for i, (lv, name, sys, user, col) in enumerate(layers):
        x = Inches(0.6) + Inches(i * 4.15)
        round_rect(s, x, Inches(1.5), Inches(3.95), Inches(4.6), C_LIGHT_CARD)
        shape = rect(s, x, Inches(1.5), Inches(3.95), Inches(0.12), col)
        add_textbox(s, x + Inches(0.25), Inches(1.85), Inches(3.4), Inches(0.35),
                    lv, size=12, bold=True, color=col)
        add_textbox(s, x + Inches(0.25), Inches(2.25), Inches(3.4), Inches(0.4),
                    name, size=20, bold=True, color=C_LIGHT_TEXT)
        add_textbox(s, x + Inches(0.25), Inches(2.9), Inches(3.4), Inches(1.0),
                    sys, size=13, color=C_LIGHT_MUTED)
        add_textbox(s, x + Inches(0.25), Inches(4.2), Inches(3.4), Inches(0.3),
                    "주 사용자", size=11, bold=True, color=C_LIGHT_ACCENT)
        add_textbox(s, x + Inches(0.25), Inches(4.55), Inches(3.4), Inches(0.8),
                    user, size=13, color=C_LIGHT_TEXT)
    slides_meta.append(("scope", False))

    # ---------- 4 KR Market ----------
    s = prs.slides.add_slide(blank)
    light_bg(s)
    section_title(s, "한국 시장 — 제품 맵", "MES · ERP · APS 주요 웹 업무시스템", dark=False)
    kr_cols = [
        ("MES / MOM", [
            "미라콤 Nexplant MESplus — 451+ 코어화면, 통합 UX",
            "삼성SDS Nexplant MES — 반도체·디스플레이·전지",
            "LG CNS Factova MES — AI 품질·설비 통합",
            "영림원 K-System MES — ERP+MES+POP",
            "더존 MES 10 — ERP10 확장, 현장 결재",
            "이맥스 Frame7 — Kiosk·단일 DB 통합",
        ]),
        ("ERP", [
            "더존 Amaranth 10 — Portal/TANGO, 통합검색",
            "더존 ERP 10 / DEWS — 구축형 확장 플랫폼",
            "영림원 K-System Ace / Ace I&I — 프로세스 메뉴",
            "시스템에버 등 SaaS — 중소 클라우드",
        ]),
        ("APS", [
            "VMS MOZART — AI·디지털트윈 시뮬레이션",
            "Asprova — 고속 Gantt (국내 다수 도입)",
            "ERP 내장 APS — 영림원·이맥스 등 패키지형",
        ]),
    ]
    for i, (title, lines) in enumerate(kr_cols):
        x = Inches(0.5) + Inches(i * 4.2)
        card_with_text(s, x, Inches(1.4), Inches(4.0), Inches(5.2),
                       title, lines, dark=False)
    slides_meta.append(("kr", False))

    # ---------- 5 KR Insights ----------
    s = prs.slides.add_slide(blank)
    dark_bg(s)
    section_title(s, "한국 시장 UI 관찰", "시안 벤치마킹 시 유의점")
    insights = [
        ("ERP↔MES 패키지 수렴", "중견 시장은 ERP 그리드 톤이 MES까지 이어지는 경우가 많음. 현장성 약화 위험."),
        ("포털·결재·검색이 L4 차별점", "Amaranth의 포틀릿·통합검색·전자결재가 한국형 업무 습관의 핵심."),
        ("대기업 MES는 SI 커스텀 비중 高", "공개 UI 시안이 적고, 업종 템플릿+커스터마이징이 기본."),
        ("POP/Kiosk는 사무 UI와 분리", "대형 버튼·터치·실적 입력 — Design System 공유하되 밀도는 분리 필요."),
        ("APS는 전문 솔루션이 UI 주도", "MOZART·Asprova의 Gantt/시뮬이 벤치마크 중심."),
    ]
    for i, (t, d) in enumerate(insights):
        y = Inches(1.4) + Inches(i * 1.0)
        round_rect(s, Inches(0.6), y, Inches(12.1), Inches(0.88), C_CARD)
        rect(s, Inches(0.6), y, Inches(0.1), Inches(0.88), C_ACCENT if i % 2 == 0 else C_ACCENT2)
        add_textbox(s, Inches(1.0), y + Inches(0.12), Inches(11.4), Inches(0.3),
                    t, size=15, bold=True, color=C_TEXT)
        add_textbox(s, Inches(1.0), y + Inches(0.45), Inches(11.4), Inches(0.35),
                    d, size=12, color=C_MUTED)
    slides_meta.append(("kr-insight", True))

    # ---------- 6 Global Market ----------
    s = prs.slides.add_slide(blank)
    light_bg(s)
    section_title(s, "글로벌 시장 — 제품 맵", "Design System · Role-based · Composable UI", dark=False)
    gl_cols = [
        ("MES / MOM", [
            "Siemens Opcenter — Mendix 개인화 Shopfloor",
            "Rockwell FactoryTalk — Plant model OEE",
            "AVEVA MES + OMI — Web Portal·위젯",
            "Plex / Infor / Critical — SaaS·하이테크",
        ]),
        ("ERP", [
            "SAP S/4 + Fiori Horizon — Floorplan 표준",
            "Oracle Cloud — Redwood UX",
            "Infor CloudSuite — 산업 템플릿·AI",
            "Dynamics 365 — Fluent + Power Platform",
        ]),
        ("APS", [
            "SAP PP/DS Scheduling Board — Gantt2.0",
            "Opcenter APS — Planning Board",
            "DELMIA Ortems — What-if·VIC",
            "PlanetTogether / Aspen Schedule Explorer",
        ]),
    ]
    for i, (title, lines) in enumerate(gl_cols):
        x = Inches(0.5) + Inches(i * 4.2)
        card_with_text(s, x, Inches(1.4), Inches(4.0), Inches(5.2),
                       title, lines, dark=False)
    slides_meta.append(("global", False))

    # ---------- 7 Global DNA ----------
    s = prs.slides.add_slide(blank)
    dark_bg(s)
    section_title(s, "글로벌 UI 공통 DNA", "2025–2026 제조 업무시스템 트렌드")
    dna = [
        ("Design System", "Fiori / Redwood / Fluent / Mendix\n일관 컴포넌트·토큰"),
        ("Role · Task", "전체 메뉴 노출 →\n역할/태스크 앱 분해"),
        ("Exception-first", "정상은 조용히,\n이탈·알람만 강조"),
        ("Composable", "Low-code 위젯 조립\n현장 맞춤 확장"),
        ("Digital Thread", "ERP↔MES↔APS\n컨텍스트 deep link"),
        ("AI Embedded", "별도 리포트가 아니라\n워크플로 내 추천+설명"),
    ]
    for i, (t, d) in enumerate(dna):
        col = i % 3
        row = i // 3
        x = Inches(0.55) + Inches(col * 4.2)
        y = Inches(1.5) + Inches(row * 2.5)
        round_rect(s, x, y, Inches(4.0), Inches(2.2), C_CARD)
        add_textbox(s, x + Inches(0.25), y + Inches(0.35), Inches(3.5), Inches(0.4),
                    t, size=18, bold=True, color=C_ACCENT)
        add_textbox(s, x + Inches(0.25), y + Inches(0.95), Inches(3.5), Inches(1.0),
                    d, size=13, color=C_MUTED)
    slides_meta.append(("dna", True))

    # ---------- 8 Taxonomy ----------
    s = prs.slides.add_slide(blank)
    light_bg(s)
    section_title(s, "화면 유형 Taxonomy — 12종", "infinitPINN 시안 카탈로그 기준 프레임", dark=False)
    types = [
        ("S01", "역할 포털/홈"), ("S02", "Command Center"),
        ("S03", "라인 현황판"), ("S04", "작업지시 실행"),
        ("S05", "트랜잭션 그리드"), ("S06", "Object 상세"),
        ("S07", "APS Gantt"), ("S08", "What-if 시나리오"),
        ("S09", "품질·NCR"), ("S10", "자재·Trace"),
        ("S11", "설비·보전"), ("S12", "리포트·분석"),
    ]
    for i, (sid, name) in enumerate(types):
        col = i % 4
        row = i // 4
        x = Inches(0.55) + Inches(col * 3.15)
        y = Inches(1.5) + Inches(row * 1.7)
        round_rect(s, x, y, Inches(3.0), Inches(1.45), C_LIGHT_CARD)
        rect(s, x, y, Inches(0.1), Inches(1.45), C_LIGHT_ACCENT)
        add_textbox(s, x + Inches(0.3), y + Inches(0.35), Inches(2.5), Inches(0.3),
                    sid, size=12, bold=True, color=C_LIGHT_ACCENT)
        add_textbox(s, x + Inches(0.3), y + Inches(0.7), Inches(2.5), Inches(0.4),
                    name, size=16, bold=True, color=C_LIGHT_TEXT)
    slides_meta.append(("taxonomy", False))

    # Screen type detail slides
    screen_details = [
        ("S01 역할 포털 / 홈", "L4·L3 · 전 역할",
         ["타일·포틀릿·카드 그리드", "통합검색·알림·미결 업무 우선", "개인화 배치, 앱 런치패드"],
         ["Amaranth Portal / Tango View", "SAP Fiori Launchpad", "Opcenter Mendix Portal"],
         "메뉴 트리가 아닌 Persona × 오늘 할 일 진입점"),
        ("S02 Command Center", "L4·L3 · 경영진·공장장",
         ["KPI 카드 + 트렌드 + 예외 리스트", "Overview → Drill-down", "멀티사이트 실시간"],
         ["Rockwell OEE Dashboard", "AVEVA 위젯 대시보드", "SAP Analytics 임베드"],
         "생산·품질·스케줄 KPI 통합, 장식용 3D 지양"),
        ("S03 라인/공정 현황판 (Andon)", "L3 · 감독·현장",
         ["대형 디스플레이·원거리 가독", "상태 블록, 최소 인터랙션", "ISA-101: 정상 중립·이상만 색"],
         ["MES Andon 보드", "High-Performance HMI", "POP 현황판"],
         "사무 UI와 시각 언어 분리 (관제 모드)"),
        ("S04 작업지시 실행 (Shopfloor)", "L3 · 오퍼레이터",
         ["현재 작업 1개 중심·대형 CTA", "단계 위저드, 바코드/서명", "도면·문서 첨부, 터치 최적화"],
         ["Opcenter Shopfloor UI", "국내 POP/Kiosk", "AVEVA MES Web Portal"],
         "글러브·조명 환경을 가정한 물리 UX"),
        ("S05 마스터/트랜잭션 그리드", "L4 · 사무·관리",
         ["필터바 + 테이블 + 툴바", "엑셀형 편집·일괄처리", "컬럼 개인화·유효성 검증"],
         ["SAP Fiori List Report", "더존/영림원 ERP 그리드", "Oracle/NetSuite 그리드"],
         "밀도는 유지, 빈상태·오류 메시지 품질로 차별"),
        ("S06 Object 상세 페이지", "L4·L3 · 사무·품질",
         ["헤더 KPI + 탭 구조", "이력·문서·관계 탭", "오더↔Lot↔설비 deep link"],
         ["Fiori Object Page", "MES Lot 상세", "Windchill Navigate형"],
         "Digital Thread의 허브 화면으로 설계"),
        ("S07 APS 스케줄 Gantt", "L3 · 스케줄러",
         ["좌 리소스 테이블 + 우 타임라인", "DnD, 제약 위반 하이라이트", "Side panel, pegging 관계선"],
         ["SAP PP/DS Scheduling Board", "PlanetTogether / Asprova", "VMS MOZART"],
         "팝업 과다 지양 → Side panel + 범례·패턴"),
        ("S08 What-if / 시나리오 비교", "L3 · 플래너",
         ["A/B 병치 또는 KPI 델타", "시나리오 저장·적용·롤백", "최적화 슬라이더"],
         ["DELMIA Ortems", "PlanetTogether What-if", "MOZART 시뮬레이션"],
         "AI 추천 시 ‘왜 이 계획인가’ 설명 패널"),
        ("S09 품질검사 · NCR", "L3 · 품질",
         ["체크리스트 + Pass/Fail CTA", "도면/이미지 오버레이", "부적합 워크플로·감사추적"],
         ["Opcenter Quality", "eDHR 패턴", "Factova AI 검사"],
         "전자기록·규제 대응 UX 내재화"),
        ("S10 자재 · 재고 · Trace", "L3·L4 · 물류·품질",
         ["재고 그리드 + Genealogy", "Lot/Serial 홀드·릴리즈", "FIFO·리콜 역추적"],
         ["MES/WMS Trace 화면", "Genealogy 트리/그래프"],
         "리콜 시나리오 원클릭 역추적"),
        ("S11 설비 · 보전", "L3 · 설비",
         ["가동상태·알람·사유코드", "PM 일정·예지보전 점수", "보전 티켓 ↔ MES 실적"],
         ["EAM 연동 UI", "Factova Control", "Rockwell Performance"],
         "MES와 보전의 양방향 링크"),
        ("S12 리포트 · 분석", "L4·L3 · 분석가",
         ["위젯 캔버스 / 고정 리포트", "임베디드 차트", "Excel·PDF 내보내기"],
         ["Embedded Analytics", "Schneider/AVEVA 위젯"],
         "트랜잭션 화면 인라인 인사이트"),
    ]

    for title, meta, feats, benches, tip in screen_details:
        s = prs.slides.add_slide(blank)
        dark_bg(s)
        section_title(s, title, meta)
        # features
        round_rect(s, Inches(0.5), Inches(1.4), Inches(6.0), Inches(4.8), C_CARD)
        add_textbox(s, Inches(0.8), Inches(1.6), Inches(5.4), Inches(0.35),
                    "화면 특징", size=14, bold=True, color=C_ACCENT)
        box = s.shapes.add_textbox(Inches(0.8), Inches(2.1), Inches(5.4), Inches(3.8))
        tf = box.text_frame
        tf.word_wrap = True
        first = True
        for f in feats:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(12)
            run = p.add_run()
            run.text = "▸  " + f
            set_run(run, size=15, color=C_TEXT)
        # benchmarks
        round_rect(s, Inches(6.8), Inches(1.4), Inches(5.9), Inches(2.8), C_CARD)
        add_textbox(s, Inches(7.1), Inches(1.6), Inches(5.4), Inches(0.35),
                    "벤치마크 레퍼런스", size=14, bold=True, color=C_ACCENT2)
        box = s.shapes.add_textbox(Inches(7.1), Inches(2.1), Inches(5.4), Inches(1.9))
        tf = box.text_frame
        tf.word_wrap = True
        first = True
        for b in benches:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(8)
            run = p.add_run()
            run.text = "•  " + b
            set_run(run, size=13, color=C_MUTED)
        # tip
        round_rect(s, Inches(6.8), Inches(4.5), Inches(5.9), Inches(1.7), C_SURFACE)
        add_textbox(s, Inches(7.1), Inches(4.7), Inches(5.4), Inches(0.3),
                    "소르테크 시사점", size=12, bold=True, color=C_ACCENT)
        add_textbox(s, Inches(7.1), Inches(5.15), Inches(5.4), Inches(0.85),
                    tip, size=14, color=C_TEXT)
        slides_meta.append(("screen", True))

    # ---------- Matrix ----------
    s = prs.slides.add_slide(blank)
    light_bg(s)
    section_title(s, "벤치마크 매트릭스 요약", "강점 레퍼런스 vs 피해야 할 패턴", dark=False)
    headers = ["화면", "한국 강점", "글로벌 강점", "Avoid"]
    rows = [
        ["S01 포털", "Amaranth 개인화", "Fiori Launchpad", "깊은 메뉴만"],
        ["S02 관제", "대기업 MES 현황", "Rockwell/AVEVA", "네온 장식 대시보드"],
        ["S03 Andon", "POP/Kiosk", "ISA-101 HMI", "작은 폰트·차트 과다"],
        ["S04 실행", "현장 Kiosk", "Opcenter Shopfloor", "사무폼 그대로 터치"],
        ["S05 그리드", "더존/영림원", "Fiori List Report", "검증없는 엑셀복제"],
        ["S07 Gantt", "MOZART/Asprova", "SAP PSB/Planet", "팝업·범례 부재"],
        ["S08 What-if", "MOZART 시뮬", "Ortems/Planet", "근거없는 AI 추천"],
    ]
    # header
    y0 = Inches(1.35)
    xs = [Inches(0.45), Inches(2.3), Inches(5.3), Inches(8.8)]
    widths = [Inches(1.75), Inches(2.9), Inches(3.4), Inches(3.9)]
    for i, h in enumerate(headers):
        round_rect(s, xs[i], y0, widths[i], Inches(0.45), C_LIGHT_ACCENT)
        add_textbox(s, xs[i] + Inches(0.1), y0 + Inches(0.08), widths[i] - Inches(0.15),
                    Inches(0.3), h, size=12, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    for r, row in enumerate(rows):
        y = y0 + Inches(0.55) + Inches(r * 0.7)
        bg = C_LIGHT_CARD if r % 2 == 0 else RgbColor(0xEA, 0xF0, 0xF4)
        for i, cell in enumerate(row):
            round_rect(s, xs[i], y, widths[i], Inches(0.6), bg)
            add_textbox(s, xs[i] + Inches(0.1), y + Inches(0.15), widths[i] - Inches(0.15),
                        Inches(0.35), cell, size=11, color=C_LIGHT_TEXT, align=PP_ALIGN.CENTER)
    slides_meta.append(("matrix", False))

    # ---------- Differentiation ----------
    s = prs.slides.add_slide(blank)
    dark_bg(s)
    section_title(s, "소르테크 · infinitPINN 차별화 방향", "시안 단계 설계 원칙 6가지")
    diffs = [
        ("01", "One Design System", "L3~L4 톤은 공유, 밀도·터치·정보계층은 분리"),
        ("02", "Persona × Task × Device", "모듈명 대신 역할·과업·단말 매트릭스"),
        ("03", "Exception & Action", "데이터 전시 → 지금 해야 할 일"),
        ("04", "Digital Thread Hub", "Object Page로 오더·Lot·설비·스케줄 연결"),
        ("05", "Explainable AI/PINN", "추천은 Side panel 근거형 UX"),
        ("06", "Korea × Global", "결재·엑셀·검색 + Design System 품질"),
    ]
    for i, (num, t, d) in enumerate(diffs):
        col = i % 3
        row = i // 3
        x = Inches(0.5) + Inches(col * 4.2)
        y = Inches(1.45) + Inches(row * 2.55)
        round_rect(s, x, y, Inches(4.0), Inches(2.3), C_CARD)
        add_textbox(s, x + Inches(0.25), y + Inches(0.3), Inches(3.5), Inches(0.35),
                    num, size=20, bold=True, color=C_ACCENT)
        add_textbox(s, x + Inches(0.25), y + Inches(0.8), Inches(3.5), Inches(0.4),
                    t, size=16, bold=True, color=C_TEXT)
        add_textbox(s, x + Inches(0.25), y + Inches(1.35), Inches(3.5), Inches(0.7),
                    d, size=13, color=C_MUTED)
    slides_meta.append(("diff", True))

    # ---------- Next steps ----------
    s = prs.slides.add_slide(blank)
    light_bg(s)
    section_title(s, "다음 단계 제안", "시안 제작 로드맵", dark=False)
    steps = [
        ("1", "Persona 확정", "공장장 / 스케줄러 / 오퍼레이터 / 품질 / 경영지원"),
        ("2", "우선 화면 6종", "S01, S02, S04, S05, S07, S09 시안 우선"),
        ("3", "Design Token", "색·타이포·상태색(정상/주의/위험/진행) 정의"),
        ("4", "와이어→하이파이", "정보구조 → 인터랙션 → 비주얼 브랜드"),
        ("5", "현장 검증", "터치거리·조명·교대 시나리오로 사용성 점검"),
    ]
    for i, (n, t, d) in enumerate(steps):
        y = Inches(1.45) + Inches(i * 1.0)
        round_rect(s, Inches(0.6), y, Inches(12.1), Inches(0.85), C_LIGHT_CARD)
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.9), y + Inches(0.18),
                                  Inches(0.5), Inches(0.5))
        fill_shape(circ, C_LIGHT_ACCENT)
        add_textbox(s, Inches(0.9), y + Inches(0.25), Inches(0.5), Inches(0.4),
                    n, size=16, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        add_textbox(s, Inches(1.7), y + Inches(0.15), Inches(4), Inches(0.35),
                    t, size=16, bold=True, color=C_LIGHT_TEXT)
        add_textbox(s, Inches(5.5), y + Inches(0.25), Inches(6.8), Inches(0.4),
                    d, size=13, color=C_LIGHT_MUTED)
    slides_meta.append(("next", False))

    # ---------- Closing ----------
    s = prs.slides.add_slide(blank)
    dark_bg(s)
    rect(s, 0, 0, Inches(0.18), SLIDE_H, C_ACCENT)
    add_textbox(s, Inches(0.7), Inches(2.4), Inches(12), Inches(0.5),
                "자료 위치", size=14, bold=True, color=C_ACCENT)
    add_textbox(s, Inches(0.7), Inches(3.0), Inches(12), Inches(1.5),
                "research/layer34-web-uisystem-benchmark-kr-global.md\n"
                "research/manufacturing-business-system-ui-report.md\n"
                "deliverables/SorTech_Layer34_UI_Benchmark.pptx",
                size=18, color=C_TEXT)
    add_textbox(s, Inches(0.7), Inches(5.2), Inches(12), Inches(0.4),
                "SorTech · infinitPINN  |  Layer 3~4 UI/UX Benchmark",
                size=13, color=C_MUTED)
    slides_meta.append(("close", True))

    # footers
    total = len(prs.slides)
    for idx, slide in enumerate(prs.slides):
        dark = slides_meta[idx][1]
        if idx == 0 or idx == total - 1:
            continue
        footer(slide, idx + 1, total, dark=dark)

    out = "/workspace/deliverables/SorTech_Layer34_UI_Benchmark.pptx"
    prs.save(out)
    # also copy to artifacts
    import shutil
    art = "/opt/cursor/artifacts/SorTech_Layer34_UI_Benchmark.pptx"
    shutil.copy(out, art)
    print(f"Saved: {out}")
    print(f"Slides: {total}")
    return out, total


if __name__ == "__main__":
    build()
